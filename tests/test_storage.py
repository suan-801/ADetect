"""저장 공간 표시·정리, 보유 상태 3단계, 01_수집정보, 프로젝트 단위 삭제."""
import os
import time
from io import BytesIO

import pytest
from openpyxl import load_workbook

INPUTS = {"brand_name": "한샘", "campaign": "", "category": "", "competitors": [],
          "brand_keywords": ["한샘"], "general_keywords": ["인테리어"],
          "sources": {"한샘": {"homepage": "https://hanssem.example/landing?utm_source=x&api_key=SECRET123", "instagram": "", "youtube": "", "meta_page": ""}},
          "paid_enabled": True, "save_images": True, "save_ad_assets": True}


def _run(pid, source, records, status="완료", created=None, **extra):
    from database.db import save_function_run, get_conn
    rid = save_function_run(pid, source, status, {"schema_version": 3, "source": source, "status": status, "records": records, "parts": {}, "inputs": INPUTS, **extra})
    if created:
        with get_conn() as conn:
            conn.execute("UPDATE function_run SET created_at=? WHERE id=?", (created, rid))
    return rid


def _age(path, days):
    old = time.time() - days * 86400
    os.utime(path, (old, old))


def test_usage_summary_levels(monkeypatch):
    from core import retention
    monkeypatch.setattr(retention, "usage", lambda: {"files": 85, "db": 10, "limits": {"files": 100, "db": 100}})
    assert retention.summary()["level"] == "warn"
    monkeypatch.setattr(retention, "usage", lambda: {"files": 10, "db": 95, "limits": {"files": 100, "db": 100}})
    state = retention.summary()
    assert state["level"] == "critical" and state["total_limit"] == 200
    monkeypatch.setattr(retention, "usage", lambda: {"files": 1, "db": 1, "limits": {"files": 100, "db": 100}})
    assert retention.summary()["level"] == "ok"


def test_capacity_error_is_detectable(monkeypatch):
    from core import retention
    from core.collection import error_message
    monkeypatch.setenv("ADETECT_STORAGE_MAX_BYTES", "1")
    with pytest.raises(retention.CapacityError) as caught:
        retention.ensure_capacity(10)
    assert retention.is_capacity_error(caught.value)
    assert error_message(caught.value) == retention.CAPACITY_MESSAGE


def test_cleanup_preview_matches_apply_and_protects_latest_pinned_and_tmp():
    from core import projects, retention
    from core.evidence_store import save_bytes, root
    pid = projects.create("한샘 조사", INPUTS)
    latest_asset = save_bytes(b"latest", "png", "https://example.com/a")
    orphan = save_bytes(b"orphan", "png", "https://example.com/b")
    kept = save_bytes(b"pinned", "png", "https://example.com/c")
    legacy = root() / ("f" * 64 + ".png")
    legacy.write_bytes(b"legacy file without metadata")
    tmp = root() / ("e" * 64 + ".tmp")
    tmp.write_bytes(b"in progress")
    for name in (latest_asset["filename"], orphan["filename"], kept["filename"], legacy.name, tmp.name):
        _age(root() / name, 60)
    from database.db import get_conn
    with get_conn() as conn:
        conn.execute("UPDATE stored_blob SET created=?", (time.time() - 60 * 86400,))
    _run(pid, "website", [{"id": "a", "kind": "공식 페이지", "brand": "한샘", "source_url": "https://example.com/a", "text": "x", "assets": [latest_asset]}], created="2020-01-01T00:00:00+00:00")
    retention.pin([kept["filename"]], True)

    preview = retention.cleanup(dry_run=True)
    names = {c["파일명"] for c in preview}
    assert orphan["filename"] in names and legacy.name in names
    assert latest_asset["filename"] not in names, "최근 성공 결과의 원본은 보호"
    assert kept["filename"] not in names, "pin 원본은 제외"
    assert tmp.name not in names, "작성 중 파일은 제외"
    assert all({"종류", "파일명", "예상 용량"} <= set(c) for c in preview)
    assert (root() / orphan["filename"]).exists(), "미리보기는 삭제하지 않는다"

    # 사용자가 본 후보 중 일부만 승인하면 그것만 삭제한다.
    approved = {retention.candidate_key(c) for c in preview if c["파일명"] == orphan["filename"]}
    done = retention.cleanup(dry_run=False, only=approved)
    assert [c["파일명"] for c in done] == [orphan["filename"]]
    assert not (root() / orphan["filename"]).exists()
    assert legacy.exists() and (root() / latest_asset["filename"]).exists() and (root() / kept["filename"]).exists()


def test_cleanup_skips_during_collection():
    from core import projects, retention
    from database.db import get_conn
    pid = projects.create("수집 중", INPUTS)
    with get_conn() as conn:
        projects.schema(conn)
        conn.execute("INSERT INTO project_task VALUES('t',?,'trend','수집 중','{}',?,'x',NULL)", (pid, projects.now()))
    assert retention.cleanup(dry_run=True)[0]["종류"] == "안내"


def test_availability_has_three_states_and_lists_missing():
    from core import retention
    from core.evidence_store import save_bytes, root
    asset = save_bytes(b"img", "png", "https://example.com/img")
    record = {"id": "r1", "kind": "광고", "brand": "A", "source_url": "https://example.com", "assets": [asset]}
    result = {"status": "완료", "records": [record]}
    assert retention.availability(result) == "결과·원본 모두 있음"
    (root() / asset["filename"]).unlink()
    state, missing = retention.availability_detail(result)
    assert state == "결과만 있음 (일부 원본 없음)"
    assert missing[0]["파일명"] == asset["filename"] and missing[0]["원문 링크"] == "https://example.com/img"
    assert retention.availability({"expired": True, "status": "완료"}) == "이력만 있음"
    assert retention.availability({"status": "실패", "records": []}) == "이력만 있음"


def test_collection_info_rows_are_readable_and_shared_by_html_and_excel():
    from core import projects
    from core.exporters import facts_report
    pid = projects.create("한샘 조사", INPUTS)
    _run(pid, "trend", [], created="2026-09-01T00:00:00+00:00", cache_policy="24시간 캐시 허용")
    _run(pid, "news", [{"id": "n", "kind": "뉴스", "brand": "한샘", "source_url": "https://news.example/a", "text": "기사", "collected_at": "now"}], created="2026-09-20T00:00:00+00:00")
    project = projects.load(pid)
    snapshots = [h for h in projects.histories(pid) if projects.available(h)]
    next(h for h in snapshots if h["source"] == "trend")["result"]["parts"] ={"trend:한샘": {"label": "t", "state": "완료", "series": {"keyword": "한샘", "start": "2023-09-01", "end": "2026-08-31", "rows": []}}}
    result = projects.selection_result(project, snapshots, [])
    assert result["inputs"]["date_mixed"] is True
    rows = facts_report.collection_info_v4(result)
    values = {r["항목"]: r["값"] for r in rows}
    assert values["프로젝트명"] == "한샘 조사" and values["자사 · 한샘 검색어 묶음"] == "한샘"
    assert values["조사 대상"] == "미입력" and values["한샘 Instagram"] == "미입력"
    assert values["시장 관심 검색어"] == "인테리어"
    assert "api_key" not in values["한샘 공식 URL 1"] and "SECRET123" not in str(rows)
    assert values["한샘 공식 URL 1"].startswith("https://hanssem.example/landing")
    assert "수집 시각 · 관련 뉴스" in values and "수집 시각 · 검색 관심도 추이" in values
    assert "24시간 캐시 허용" in values["캐시 사용 여부"] and "기록 없음" in values["캐시 사용 여부"]
    assert values["SAMPLE 여부"] == "실수집 자료" and values["유료 기능"] == "ON"
    assert "날짜 혼합 안내" in values
    assert not any(isinstance(v, (dict, list)) or str(v).startswith("{") for v in values.values()), "JSON 덤프 금지"

    html = facts_report.report_html(project, result)
    assert '<a href="https://hanssem.example/landing' in html and "SECRET123" not in html
    sheet = load_workbook(BytesIO(facts_report.excel(result)))["17_수집조건"]
    excel_rows = [(r[0].value, r[1].value) for r in sheet.iter_rows(min_row=2)]
    assert excel_rows == [(r["항목"], r["값"]) for r in rows], "HTML·Excel 공통 표"
    link = next(r[1] for r in sheet.iter_rows(min_row=2) if r[0].value == "한샘 공식 URL 1")
    assert link.hyperlink is not None


def test_project_delete_requires_exact_name_and_is_project_scoped():
    from core import projects
    keep = projects.create("유지", INPUTS)
    drop = projects.create("삭제 대상", INPUTS)
    _run(keep, "news", [])
    _run(drop, "news", [], sample_sources=["SAMPLE"])
    kinds = {r["name"]: r["kind"] for r in projects.overview()}
    assert kinds == {"유지": "실수집 자료 포함", "삭제 대상": "SAMPLE 자료만 있음"}
    assert projects.delete_impact(drop)["수집 실행"] == 1
    with pytest.raises(ValueError):
        projects.delete(drop, "삭제")
    projects.delete(drop, "삭제 대상")
    assert [r["name"] for r in projects.overview()] == ["유지"]
    assert len(projects.histories(keep)) == 1


def test_db_backup_copies_database(tmp_path, monkeypatch):
    import sqlite3
    from core import projects, db_backup
    monkeypatch.setenv("ADETECT_BACKUP_DIR", str(tmp_path / "backups"))
    projects.create("백업 확인", INPUTS)
    path = db_backup.backup_db()
    with sqlite3.connect(path) as conn:
        assert conn.execute("SELECT name FROM project").fetchall() == [("백업 확인",)]
    assert db_backup.backups() == [path]
