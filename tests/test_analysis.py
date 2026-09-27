from io import BytesIO
from threading import Event
from openpyxl import load_workbook
from core.analyzers.market_analyzer import seasonality, run_market_analysis
from core.analyzers.insight_synthesizer import run_synthesis
from core.exporters.report_builder import build_report_excel, build_report_html
from core.runtime import persist_result, restore_results
from database.db import create_session, list_function_runs


def test_seasonality_zero_and_missing():
    assert seasonality([])["weekend_gap_pct"] is None
    result = seasonality([{"date":"2026-09-21","search_index":10},{"date":"2026-09-26","search_index":20}])
    assert result["weekend_gap_pct"] == 100
    assert result["monthly"][0]["search_index"] == 15


def test_market_news_failure_preserves_trend(monkeypatch):
    import core.analyzers.market_analyzer as m
    monkeypatch.setattr(m,"get_search_volume_trend",lambda *a,**k:[{"date":"2026-09-21","search_index":10}])
    def fail(*a,**k):
        raise RuntimeError("failed")
    monkeypatch.setattr(m,"get_news",fail)
    result = run_market_analysis("테스트")
    assert result["status"] == "부분 실패"
    assert result["trend"]
    assert result["upcoming_changes"] == []


def test_synthesis_does_not_divide_missing_brand_or_zero():
    session={"brand_name":"A","competitors":["B"],"creative_status":"부분 실패",
             "creative_result":{"own":{"brand":"A","ad_count":10},"competitors":[]}}
    result = run_synthesis(session)
    assert all(r["share_of_ads"] is None for r in result["sov"])
    session["creative_result"]["competitors"]=[{"brand":"B","ad_count":30}]
    result = run_synthesis(session)
    assert [r["share_of_ads"] for r in result["sov"]] == [25,75]
    assert result["target_insight"]["status"] == "insufficient_data"


def test_search_anchor_shared_groups(monkeypatch):
    import core.scrapers.naver_api as api
    monkeypatch.setattr(api.settings,"NAVER_DATALAB_MOCK",False)
    calls=[]
    class Response:
        def raise_for_status(self): pass
        def json(self):
            groups=calls[-1]["keywordGroups"]
            factor=1 if len(calls)==1 else 2
            return {"results":[{"title":g["groupName"],"data":[{"period":"2026-09-01","ratio":10*factor}]} for g in groups]}
    def post(*a,**kw):
        calls.append(kw["json"])
        return Response()
    monkeypatch.setattr(api.requests,"post",post)
    result=api.get_comparable_brand_trends("A",["B","C","D","E","F"],{"A":["A","AA"]})
    assert len(calls)==2
    assert all(c["keywordGroups"][0]["groupName"]=="A" for c in calls)
    assert calls[0]["keywordGroups"][0]["keywords"]==["A","AA"]
    assert result["F"]["relative_trend"][0]["search_index"]==10


def test_exports_escape_html_and_formula():
    result={"status":"완료","sample_sources":["trend"],"trend":[{"date":"=1+1","search_index":4}],
            "reference_sources":[{"name":"<script>alert(1)</script>"}]}
    workbook=load_workbook(BytesIO(build_report_excel("market",result)))
    assert workbook["02_Market_Trend"]["A2"].data_type=="s"
    html=build_report_html({"brand_name":"<bad>"},"market",result)
    assert "<script>" not in html
    assert "&lt;bad&gt;" in html and "SAMPLE" in html
    assert not any(n.startswith("05") for n in workbook.sheetnames)


def test_runs_restore_latest_and_invalidate_synthesis():
    session={"id":create_session("A","cat",[])}
    persist_result(session,"market",{"status":"완료","trend":[]})
    persist_result(session,"synthesis",{"status":"완료","input_signature":"stale"})
    restored=restore_results({"id":session["id"]})
    assert restored["market_status"]=="완료"
    assert restored["synthesis_status"]=="미실행"
    assert len(list_function_runs(session["id"]))==2


def test_job_queue_cancellation_and_cache():
    from core.jobs import submit,get_job,cancel,checkpoint
    one={"id":create_session("A","cat",[])}
    two={"id":create_session("B","cat",[])}
    entered,release=Event(),Event()
    def slow(snapshot):
        entered.set()
        release.wait(5)
        checkpoint()
        return {"status":"완료","value":1}
    jid=submit(one,"market",slow,{"test":"one"})
    assert entered.wait(5)
    called=[]
    other=submit(two,"market",lambda s:called.append(1) or {"status":"완료"},{"test":"two"})
    cancel(other)
    release.set()
    get_job(jid)["future"].result(10)
    get_job(other)["future"].result(10)
    assert not called
    assert get_job(other)["result"]["status"]=="취소"
    cached=submit(one,"market",lambda s:(_ for _ in ()).throw(AssertionError()),{"test":"one"})
    get_job(cached)["future"].result(10)
    assert get_job(cached)["result"]["cache_hit"] is True


def test_brand_failures_are_isolated(monkeypatch):
    import core.analyzers.brand_analyzer as b
    def fail(*a,**kw):
        raise RuntimeError("secret-must-not-be-stored")
    monkeypatch.setattr(b,"fetch_meta_ads_detail",fail)
    r=b.run_brand_analysis("A",["B"],collect_naver_sa=False,collect_instagram=False)
    assert r["status"]=="부분 실패"
    assert len(r["competitors"])==1
    assert r["own"]["ad_count"] is None
    assert r["own"]["media_operation_matrix_row"]["naver_sa"] is None
    assert "secret-must-not-be-stored" not in str(r)


def test_artifact_expiration_preserves_history(monkeypatch):
    from core.exporters.artifact_store import save_artifact,read_artifact,list_artifacts
    monkeypatch.setenv("ADETECT_EXPORT_MAX_RUNS","1")
    session={"id":create_session("A","cat",[])}
    old=persist_result(session,"market",{"status":"완료"})
    a=save_artifact(old,"market","html","one")
    new=persist_result(session,"market",{"status":"완료"})
    b=save_artifact(new,"market","html","two")
    assert read_artifact(a) is None
    assert read_artifact(b)==b"two"
    assert list_artifacts(old)[0]["expired"]
    assert len(list_function_runs(session["id"]))==2


def test_ai_rejects_fabricated_sources_and_quotes(monkeypatch):
    import json
    from types import SimpleNamespace
    from core.analyzers import evidence, gemini_client
    monkeypatch.setattr(evidence.settings,"GEMINI_MOCK",False)
    payload={"valid":{"insight":"관측된 혜택","source":["official"],"evidence":["무료 배송"],"confidence":"low"},
             "invented":{"insight":"지어낸 해석","source":["official"],"evidence":["매출 200% 상승"],"confidence":"high"},
             "wrong_source":{"insight":"출처 위조","source":["fake"],"evidence":["무료 배송"],"confidence":"low"}}
    fake=SimpleNamespace(models=SimpleNamespace(generate_content=lambda **kw:SimpleNamespace(text=json.dumps(payload))))
    monkeypatch.setattr(gemini_client,"get_client",lambda:fake)
    result=evidence.interpret(payload.keys(),[{"source":"official","text":"오늘 무료 배송 이벤트"}])
    assert result["valid"]["insight"]=="관측된 혜택"
    assert result["invented"]["status"]=="insufficient_data"
    assert result["wrong_source"]["status"]=="insufficient_data"


def test_positioning_requires_confirmation_and_real_evidence(monkeypatch):
    from core.analyzers import positioning
    monkeypatch.setattr(positioning.settings,"GEMINI_MOCK",False)
    brand={"own":{"brand":"A","brand_website_facts":{"source_url":"a","raw_copy_snippets":["원문 A"]}},
           "competitors":[{"brand":"B","brand_website_facts":{"source_url":"b","raw_copy_snippets":["원문 B"]}}]}
    assert positioning.score_positions(brand,{"confirmed":False})["status"]=="insufficient_data"
    points=[{"brand":n,"x_axis_score":10,"y_axis_score":20,"insight":"정성 점수","source":[n.lower()],"evidence":["원문 "+n],"confidence":"low"} for n in ("A","B")]
    monkeypatch.setattr(positioning,"request_json",lambda prompt:{"points":points})
    axes={"confirmed":True,"x_axis_label":"X","y_axis_label":"Y"}
    assert positioning.score_positions(brand,axes)["status"]=="available"
    points[1]["evidence"]=["가짜 인용"]
    assert positioning.score_positions(brand,axes)["status"]=="insufficient_data"


def test_public_url_rejects_local_network():
    import pytest
    from core.scrapers.brand_site import validate_public_url
    for url in ("file:///etc/passwd","http://127.0.0.1/","http://[::1]/","https://user:pass@example.com/"):
        with pytest.raises(ValueError):
            validate_public_url(url)


def test_meta_resolution_does_not_choose_most_frequent_unknown():
    from core.scrapers.ad_library import _resolve_brand_page
    unrelated=[{"page_name":"리셀러"}]*10
    assert _resolve_brand_page(unrelated,"브랜드") is None
    assert _resolve_brand_page(unrelated+[{"page_name":"브랜드"}],"브랜드")=="브랜드"
    verified=[{"page_name":"A","page_like_count":10000},{"page_name":"B","is_verified":True,"page_like_count":20}]
    assert _resolve_brand_page(verified,"브랜드")=="B"


def test_target_zero_baseline_does_not_claim_flat():
    from core.analyzers.insight_synthesizer import build_target_insight
    result=build_target_insight({"trend":[{"search_index":0},{"search_index":10}]},{"own":{"brand":"A"}})
    assert "판단 불가" in result["insight"]
    assert "+0.0%" not in str(result)


def test_website_falls_back_and_preserves_verbatim(monkeypatch):
    import core.scrapers.brand_site as site
    def fetch(url):
        if url.endswith("detail"):
            raise ValueError("failed")
        return '<html><body><p>'+('원문입니다. '*25)+'</p><script>hidden secret</script></body></html>',url
    monkeypatch.setattr(site,"fetch_public_html",fetch)
    result=site.crawl_brand_website("A","https://example.com/detail","https://example.com")
    assert result["status"]=="available"
    assert result["fallback_used"]
    assert "hidden secret" not in str(result)
    assert result["raw_copy_snippets"][0].startswith("원문입니다.")
