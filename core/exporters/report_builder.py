"""PRD §14 기능별 HTML/Excel. 모든 외부 텍스트를 이스케이프한다."""
import html
import json
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font


def tables(function, result):
    tables = {"01_Overview": [{"function": function, "status": result.get("status"), "collected_at": result.get("collected_at"),
                               "sample_sources": result.get("sample_sources", []), "errors": result.get("errors", [])}]}
    if function == "market":
        tables.update({"02_Market_Trend": result.get("trend", []),
                       "03_Market_Seasonality": [result.get("search_seasonality", {})],
                       "04_Market_References": result.get("reference_sources", []),
                       "13_News": [{"scope": "market", **n} for n in result.get("news", [])]})
    if function == "brand":
        profiles = [result.get("own", {}), *result.get("competitors", [])]
        tables.update({"06_Brand_Profile": [{k:v for k,v in p.items() if k != "brand_context"} for p in profiles],
                       "07_Naver_Operation": [{"brand": p["brand"], "brand_search": p.get("naver_brand_search"), **r} for p in profiles for r in (p.get("naver_sa_ranking") or [{}])],
                       "08_Media_Matrix": [{"brand": p["brand"], **p.get("media_operation_matrix_row", {})} for p in profiles],
                       "13_News": [{"scope": "brand", "brand": p["brand"], **n} for p in profiles for n in p.get("brand_news", [])],
                       "14_UTM": [r for p in profiles for r in p.get("meta_utm", [])]})
    if function == "synthesis":
        tables["15_SOV"] = result.get("sov", [])
        tables["16_Positioning"] = result.get("positioning_map", {}).get("points") or [result.get("positioning_map", {})]
    insights = []
    for key, value in result.items():
        if isinstance(value, dict) and ("insight" in value or "status" in value):
            insights.append({"field": key, "classification": "REC" if key in ("recommended_angle", "one_line_summary") else "AI", **value})
    if insights:
        tables["17_AI_Insights"] = insights
    return tables


def without_images(value):
    if isinstance(value,dict):
        return {k:without_images(v) for k,v in value.items() if k != "capture_data_uri"}
    if isinstance(value,list):
        return [without_images(v) for v in value]
    return value


def scalar(value):
    if isinstance(value, (list, dict)):
        return json.dumps(without_images(value), ensure_ascii=False)
    return value


def build_report_excel(function, result):
    if result.get("schema_version") == 2:
        from core.exporters.facts_report import excel
        return excel(result)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tables(function, result).items():
        ws = wb.create_sheet(name)
        headers = list(dict.fromkeys(k for row in rows for k in row)) or ["status"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([scalar(row.get(k)) for k in headers])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # 외부 카피를 Excel 수식으로 실행하지 않는다.
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def build_report_html(session, function, result):
    if result.get("schema_version") == 2:
        from core.exporters.facts_report import report_html
        return report_html(session,result)
    esc = lambda v: html.escape(str(scalar(v) if v is not None else "not_available"))
    def render_value(value):
        if isinstance(value, dict):
            if "insight" in value:
                return '<p>'+esc(value["insight"])+"</p><small>Confidence: "+esc(value.get("confidence"))+"</small>"+render_value({"출처":value.get("source"),"근거":value.get("evidence")})
            if value.get("status") in ("insufficient_data","not_available","not_collected","unknown"):
                return '<p>'+esc(value.get("message") or value["status"])+"</p>"
            return '<dl>'+''.join('<dt>'+esc(k)+'</dt><dd>'+render_value(v)+'</dd>' for k,v in value.items() if k not in ("capture_data_uri","brand_context"))+'</dl>'
        if isinstance(value,list):
            return '<ul>'+''.join('<li>'+render_value(v)+'</li>' for v in value)+'</ul>' if value else '<span>없음</span>'
        if isinstance(value,str) and value.startswith(("http://","https://")):
            return '<a href="'+html.escape(value,quote=True)+'" rel="noreferrer">'+esc(value)+'</a>'
        return esc(value)
    sections = []
    for name, rows in tables(function, result).items():
        headers = list(dict.fromkeys(k for row in rows for k in row))
        body = "".join("<tr>"+"".join("<td>"+render_value(row.get(k))+"</td>" for k in headers)+"</tr>" for row in rows)
        sections.append("<section><h2>"+esc(name)+"</h2><div class=scroll><table><thead><tr>"+
                        "".join("<th>"+esc(k)+"</th>" for k in headers)+"</tr></thead><tbody>"+body+"</tbody></table></div></section>")
    if function == "market" and result.get("trend"):
        trend=result["trend"]
        maximum=max((r["search_index"] for r in trend), default=0) or 1
        points=" ".join(f"{20+i*760/max(1,len(trend)-1):.1f},{220-r['search_index']/maximum*200:.1f}" for i,r in enumerate(trend))
        sections.insert(0,'<section><h2>검색 관심도 추이 · FACT</h2><p>공통 기간 상대지수. 절대 검색 횟수가 아닙니다.</p><svg viewBox="0 0 800 250" role="img" aria-label="검색지수 추이"><polyline points="'+points+'" fill="none" stroke="#ff7858" stroke-width="2"/></svg></section>')
    if function == "brand":
        for profile in [result.get("own", {}), *result.get("competitors", [])]:
            for row in profile.get("naver_sa_ranking", []):
                image = row.get("capture_data_uri", "")
                if image.startswith("data:image/png;base64,"):
                    sections.append('<figure><img style="max-width:100%" src="'+html.escape(image,quote=True)+'"><figcaption>'+esc(profile.get("brand"))+" · "+esc(row.get("keyword"))+"</figcaption></figure>")
    if function == "synthesis" and result.get("positioning_map", {}).get("status") == "available":
        position = result["positioning_map"]
        points = "".join('<circle cx="'+str(50+p["x_axis_score"]*5)+'" cy="'+str(550-p["y_axis_score"]*5)+'" r="6" fill="#ff7858"/><text x="'+str(60+p["x_axis_score"]*5)+'" y="'+str(550-p["y_axis_score"]*5)+'" fill="white">'+esc(p["brand"])+"</text>" for p in position["points"])
        sections.insert(0,'<section><h2>포지셔닝맵 · AI 정성 해석</h2><p>'+esc(position["axes"]["x_axis_label"])+" / "+esc(position["axes"]["y_axis_label"])+
            '</p><svg viewBox="0 0 650 600" style="max-width:650px"><path d="M50 50 V550 H600" stroke="#888" fill="none"/>'+points+'</svg></section>')
    sample = "<p class=sample>SAMPLE 포함 — 실제 분석과 혼동하지 마세요: "+esc(result.get("sample_sources"))+"</p>" if result.get("sample_sources") else ""
    return ("<!doctype html><html lang=ko><meta charset=utf-8><meta name=viewport content='width=device-width,initial-scale=1'>"
            "<title>ADetect</title><style>body{background:#101010;color:#eee;font:15px system-ui;margin:40px}"
            "h1,h2{font-weight:600}a{color:#ff997f}dt{font-weight:600;color:#bbb}dd{margin:8px 0 18px}ul{padding-left:20px}small{color:#bbb}table{border-collapse:collapse;min-width:100%}td,th{text-align:left;border-bottom:1px solid #444;padding:12px;max-width:600px;white-space:pre-wrap;overflow-wrap:anywhere}"
            ".scroll{overflow:auto}.sample{color:#ff7858}section{margin:40px 0}</style><body><h1>"+esc(session["brand_name"])+" · "+esc(function)+"</h1>"+
            sample+"<p>FACT: 수집 데이터 · AI: 근거 기반 해석 · REC: 권고. 데이터 부족은 미제공합니다.</p>"+
            "".join("<p>"+esc(t)+"</p>" for t in result.get("limitations", []))+"".join(sections)+"</body></html>")
