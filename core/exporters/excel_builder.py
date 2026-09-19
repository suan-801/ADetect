"""소재분석 Excel 산출물 빌더 — PRD §14 09_Ads~11_LongRunning_Analysis.

이미지/영상은 셀에 임베드하지 않고 URL 컬럼으로만 제공한다 — 이유는 PRD §7-11-(4) 참고
(추후 필요 시 openpyxl.drawing.image.Image로 셀 삽입 확장 가능).

[제거됨] 기존 11_Appeal_Distribution 시트와 10_CreativeAnalysis.appeal_tags/
12_LongRunning_Analysis.dominant_appeal_tags 컬럼은 근거 없는 seeded_random 태깅
결과였다 — 실제 분석처럼 보이는 랜덤 값을 Excel 사용자에게 노출하지 않기 위해 제거했다
(PRD §5/§10 Guardrail). 시트 번호는 새 구조로 다시 채번했다(과거 저장된 결과가 없어
번호 호환을 지킬 필요가 없음 — save_function_run()이 어디서도 호출되지 않는 상태).
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _write_sheet(ws, headers: list[str], rows: list[list]):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        width = max(len(str(header)) + 2, 12)
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    ws.freeze_panes = "A2"


def _brand_label(b: dict) -> str:
    return b["brand"] + (" (자사)" if b["is_own"] else "")


def build_creative_excel(result: dict) -> bytes:
    own = result["own"]
    competitors = result["competitors"]
    all_brands = [own, *competitors]

    wb = Workbook()
    wb.remove(wb.active)

    ws_ads = wb.create_sheet("09_Ads")
    _write_sheet(
        ws_ads,
        ["브랜드", "platform", "publisher_platforms", "ad_id", "format", "headline", "body",
         "cta", "landing_url", "image_url", "thumbnail_url", "ad_delivery_start_time",
         "ad_running_days", "collected_at"],
        [
            [_brand_label(b), ad["platform"], ad["publisher_platforms"], ad["ad_id"], ad["format"],
             ad["headline"], ad["body"], ad["cta"], ad["landing_url"], ad.get("image_url"),
             ad.get("thumbnail_url"), ad.get("ad_delivery_start_time"), ad.get("ad_running_days"),
             ad["collected_at"]]
            for b in all_brands for ad in b["ads"]
        ],
    )

    ws_creative = wb.create_sheet("10_CreativeAnalysis")
    _write_sheet(
        ws_creative,
        ["ad_id", "브랜드", "headline", "body_copy", "cta", "format"],
        [
            [ad["ad_id"], _brand_label(b), ad["headline"], ad["body"], ad["cta"], ad["format"]]
            for b in all_brands for ad in b["ads"]
        ],
    )

    ws_long = wb.create_sheet("11_LongRunning_Analysis")
    _write_sheet(
        ws_long,
        ["브랜드", "running_days_bucket", "avg_running_days", "ad_count"],
        [
            [_brand_label(b), r["running_days_bucket"], r["avg_running_days"], r["ad_count"]]
            for b in all_brands for r in b["long_running_analysis"]
        ],
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
