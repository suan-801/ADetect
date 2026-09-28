"""팩트 중심 동일 스키마의 개별/통합 보고서와 원본 패키지."""
import html
import base64
import json
import re
import hashlib
from io import BytesIO
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook
from core.evidence_store import read_bytes
from core.exporters.artifact_store import check_secrets


MISSING="미입력"
SECRET_PARAMS=("token","key","secret","signature","password","auth")


def _safe_url(value):
    """사용자가 입력한 URL에서 인증값처럼 보이는 query 파라미터를 제거한다."""
    if not value: return MISSING
    u=urlsplit(str(value))
    if not u.scheme: return str(value)
    query=urlencode([(k,v) for k,v in parse_qsl(u.query) if not any(s in k.lower() for s in SECRET_PARAMS)])
    return urlunsplit((u.scheme,u.netloc,u.path,query,u.fragment))


def _time(value):
    """저장 시각은 UTC ISO 문자열이다. 시간대를 명시해 표시한다."""
    if not value: return MISSING
    text=str(value)[:19].replace("T"," ")
    return text+" UTC" if "+00:00" in str(value) or str(value).endswith("Z") else text


def _join(values):
    values=[str(v) for v in dict.fromkeys(values) if v not in (None,"")]
    return ", ".join(values) if values else MISSING


def official_rows(official):
    urls=[u for u in dict.fromkeys(list(official.get("official_urls") or [])+[official.get("homepage")]) if u]
    if len(urls)<=1:
        return [("공식 페이지",_safe_url(urls[0] if urls else None))]
    return [(f"공식 페이지 {i}",_safe_url(u)) for i,u in enumerate(urls,1)]


def collection_info_rows(result):
    """01_수집정보. HTML과 Excel이 같은 행을 쓴다. 입력값 화이트리스트만 사용하며 JSON을 덤프하지 않는다."""
    from core import projects
    inputs=result.get("inputs",{}) or {}
    versions=inputs.get("versions") or [{"source":result.get("source"),"collected_at":result.get("collected_at"),"inputs":inputs,
                                          "cache_policy":result.get("cache_policy"),"sample":bool(result.get("sample_sources"))}]
    snapshots=[v.get("inputs") or {} for v in versions]
    brand=_join([inputs.get("brand_name")]+[s.get("brand_name") for s in snapshots])
    brand_names=[s.get("brand_name") for s in snapshots]+[inputs.get("brand_name")]
    official={}
    for s in snapshots:
        for name,fields in (s.get("sources") or {}).items():
            if name in brand_names or not official:
                for k,v in (fields or {}).items():
                    if v: official.setdefault(k,v)
    label=lambda src: projects.SOURCES[src][0] if src in projects.SOURCES else (src or "자료")
    paid=[s.get("paid_enabled") for s in snapshots if s.get("paid_enabled") is not None]
    trend=[p["series"] for p in result.get("parts",{}).values() if p.get("series")]
    rows=[("프로젝트명",inputs.get("project") or MISSING),("브랜드명",brand),
          ("조사 대상",_join([s.get("campaign") for s in snapshots]+[inputs.get("campaign")])),
          ("브랜드/캠페인 검색어",_join([k for s in snapshots for k in s.get("brand_keywords") or []]+list(inputs.get("brand_keywords") or []))),
          ("일반 검색어",_join([k for s in snapshots for k in s.get("general_keywords") or []]+list(inputs.get("general_keywords") or []))),
          ("뉴스 검색어",_join([k for s in snapshots for k in s.get("news_keywords") or []])),
          ("검색 추이 조회 방식",_join("브랜드 표기 묶음" if s.get("keyword_mode")=="brand_group" else "개별 검색어 (이전 방식)" for s in snapshots)),
          *official_rows(official),
          ("캠페인 상세 페이지",_safe_url(official.get("detail_url"))),
          ("공식 Instagram",_safe_url(official.get("instagram"))),
          ("공식 YouTube",_safe_url(official.get("youtube"))),
          ("Meta 광고 페이지",_safe_url(official.get("meta_page"))),
          ("선택한 자료 종류",_join(label(v.get("source")) for v in versions if v.get("source")))]
    rows+=[("수집 시각 · "+label(v.get("source")),_time(v.get("collected_at"))) for v in versions]
    if inputs.get("date_mixed"):
        rows.append(("날짜 혼합 안내","선택한 자료의 수집일이 서로 다릅니다: "+", ".join(inputs.get("collection_dates",[]))+". 같은 시점 비교가 아닙니다."))
    rows+=[("검색 추이 기간",_join(f"{s.get('start','?')} ~ {s.get('end','?')}" for s in trend) if trend else "해당 없음 (검색 추이 미선택)"),
           ("캐시 사용 여부",_join(v.get("cache_policy") or "기록 없음 (이전 버전 수집)" for v in versions)),
           ("SAMPLE 여부","SAMPLE — 실제 수집 자료가 아님" if result.get("sample_sources") or any(v.get("sample") for v in versions) else "실수집 자료"),
           ("유료 기능",("ON" if all(paid) else "OFF" if not any(paid) else "자료별로 다름") if paid else "기록 없음"),
           ("보고서 생성 시각",_time(result.get("collected_at"))),("상태",result.get("status") or MISSING)]
    return [{"항목":k,"값":v} for k,v in rows]


def report_tables(result):
    parts=result.get("parts",{})
    records=result.get("records",[])
    tables = {
        "01_수집정보":collection_info_rows(result),
        "02_자료목록":[{"자료ID":r["id"],"종류":r["kind"],"브랜드":r["brand"],"확인상태":r.get("review"),"선정근거":r.get("review_reason"),"출처":r["source_url"],"원문":r["text"],"수집시각":r["collected_at"],"SAMPLE":r.get("sample",False)} for r in records],
        "03_검색추이":[{"검색어":p["series"]["keyword"],"월":r["date"],"상대지수":r["search_index"]} for p in parts.values() if p.get("series") for r in p["series"]["rows"]],
        "04_월별평균":[{"검색어":p["series"]["keyword"],**r} for p in parts.values() if p.get("seasonality") for r in p["seasonality"]["monthly"]],
        "05_연도별피크":[{"검색어":p["series"]["keyword"],**r} for p in parts.values() if p.get("seasonality") for r in p["seasonality"]["yearly"]],
        "06_검색량":[{"검색어":r["text"],"조회어":r.get("keyword"),"구분":r["kind"],"PC":r.get("pc"),"모바일":r.get("mobile"),"주의":r.get("note")} for r in records if r["kind"] in ("검색량","연관 검색어")],
        "07_광고":[{"자료ID":r["id"],"광고ID":r.get("external_id"),"문구":r["text"],"형식":r.get("format"),"CTA":r.get("cta"),"랜딩":r.get("landing_url"),"시작일":r.get("start_date"),"관측운영일수":r.get("running_days"),"지면":r.get("placements"),"표본제한":r.get("note")} for r in records if r["kind"]=="광고"],
        "08_원본파일":[{"자료ID":r["id"],**asset} for r in records for asset in r.get("assets",[])],
        "09_수집상태":[{"대상":p["label"],"상태":p["state"],"안내":p.get("message","")} for p in parts.values()],
        "10_근거요약":[result["summary"]] if result.get("summary") else [],
        "11_이전수집대비":result.get("changes",[]),
    }
    if result.get('schema_version')==3:
        from core.materials import prioritized, summary, table_row
        records = prioritized(records)
        tables["02_자료목록"] = [{k:v for k,v in row.items() if k != "자료ID"} for row in tables["02_자료목록"]]
        for name in ("07_광고", "08_원본파일", "11_이전수집대비"):
            tables[name] = [{k:v for k,v in row.items() if k != "자료ID"} for row in tables[name]]
        tables["06_검색량"] = [r for r in tables["06_검색량"] if r["구분"] == "검색량"]
        tables["12_자료별요약"] = [{"종류":kind,"요약":summary(kind,[r for r in records if r["kind"]==kind])} for kind in dict.fromkeys(r["kind"] for r in records)]
        for index, kind in enumerate(("뉴스", "Instagram", "YouTube", "연관 검색어", "홈페이지", "Instagram 게시물", "YouTube 게시물"), 13):
            tables[f"{index:02d}_{kind}"] = [{k:v for k,v in table_row(r,True).items() if k not in ("선택","다운로드")} | {"원문":r["text"]} for r in records if r["kind"]==kind]
        tables={k:v for k,v in tables.items() if v or k=='01_수집정보'}
        tables = {k:tables[k] for k in (["12_자료별요약"] if "12_자료별요약" in tables else []) + [k for k in tables if k != "12_자료별요약"]}
    return tables


def collection_info_v4(result):
    """v4 수집조건. 입력값 화이트리스트만 쓰고 API 키·민감 파라미터는 표시하지 않는다."""
    from core import projects
    inputs = result.get("inputs", {})
    rows = [("프로젝트명", inputs.get("project") or MISSING), ("조사 대상", inputs.get("campaign") or MISSING)]
    for b in inputs.get("brands", []):
        role = "자사" if b.get("role") == "own" else "경쟁사"
        rows.append((f"{role} · {b['name']} 검색어 묶음", _join(b.get("terms", []))))
        src = b.get("sources", {})
        for i, url in enumerate(official_urls_of(src), 1):
            rows.append((f"{b['name']} 공식 URL {i}", _safe_url(url)))
        for key, label in (("detail_url", "캠페인 상세 URL"), ("instagram", "Instagram"), ("youtube", "YouTube"), ("meta_page", "Meta 광고 라이브러리")):
            rows.append((f"{b['name']} {label}", _safe_url(src.get(key))))
    flt = result.get("news_filter") or {}
    rows += [("시장 관심 검색어", _join(inputs.get("market_keywords", []))), ("뉴스 검색어", _join(inputs.get("news_keywords", []))),
             ("뉴스 결과 좁히기 · 꼭 들어갈 문구", _join(flt.get("include", []))), ("뉴스 결과 좁히기 · 빼고 싶은 문구", _join(flt.get("exclude", []))),
             ("뉴스 필터로 숨긴 기사", f"{result.get('hidden_by_news_filter', 0)}건 (원본은 보관, 제목·발췌문 기준)")]
    label = lambda src: projects.SOURCES[src][0] if src in projects.SOURCES else src
    rows += [("수집 시각 · " + label(v.get("source")), _time(v.get("collected_at"))) for v in inputs.get("versions", [])]
    if inputs.get("date_mixed"):
        rows.append(("날짜 혼합 안내", "선택한 자료의 수집일이 서로 다릅니다: " + ", ".join(inputs.get("collection_dates", []))))
    rows += [("캐시 사용 여부", _join(v.get("cache_policy") or "기록 없음 (이전 버전 수집)" for v in inputs.get("versions", []))),
             ("SAMPLE 여부", "SAMPLE — 실제 수집 자료가 아님" if result.get("sample_sources") else "실수집 자료"),
             ("유료 기능", "ON" if inputs.get("paid_enabled", True) else "OFF"), ("보고서 생성 시각", _time(result.get("collected_at")))]
    return [{"항목": k, "값": v} for k, v in rows]


def official_urls_of(src):
    return [u for u in dict.fromkeys(list(src.get("official_urls") or []) + [src.get("homepage")]) if u]


def project_tables(result):
    """v4 선택 결과의 표. 요약·비교 → 상세 → 수집 조건 순서. 내부 ID는 싣지 않는다(JSON에는 유지)."""
    from core.materials import table_row, prioritized
    from core.project_view import volume_totals, search_status_rows
    from core.result_insights import all_trend_facts, news_sections
    from core import utm
    records = prioritized(result.get("records", []))
    parts = result.get("parts", {})
    brands = result.get("inputs", {}).get("brands", [])
    by_kind = lambda *kinds: [r for r in records if r["kind"] in kinds]
    clean = lambda row: {k: v for k, v in row.items() if k != "다운로드"}
    compare = next((pt["comparison"] for pt in parts.values() if pt.get("comparison")), None)
    trend = []
    if compare:
        months = sorted({r["date"][:7] for s in compare["series"] for r in s["rows"]})
        values = {(s["name"], r["date"][:7]): r["search_index"] for s in compare["series"] for r in s["rows"]}
        trend = [{"월": m, **{s["name"]: values.get((s["name"], m), "미제공") for s in compare["series"]}} for m in months]
    market = [{"검색어": pt["series"]["keyword"], "월": r["date"][:7], "상대지수(검색어별 별도 기준)": r["search_index"]}
              for k, pt in parts.items() if pt.get("source") == "trend" and pt.get("series") and not pt.get("comparison") for r in pt["series"]["rows"]]
    volume_item = {"result": {"parts": {k: pt for k, pt in parts.items() if pt.get("terms") is not None}}}
    totals = [{k: v for k, v in r.items() if k not in ("brand_id", "정확한 합계", "known_total")} for r in volume_totals(volume_item, brands + [{"id": "market", "name": "시장 관심 검색어"}])]
    ads = [{"검색어": r.get("keyword"), "영역": {"powerlink": "파워링크", "brand_search": "브랜드검색"}.get(a["area"], a["area"]), "광고 문구": a.get("text"),
            "표시 URL": a.get("display_url"), "링크": ", ".join(a.get("links") or []), "판독 방식": a.get("method")} for r in by_kind("검색 화면") for a in r.get("ads", [])]
    ads += [{"검색어": r.get("keyword"), "영역": a["area"], "광고 문구": " ".join(x for x in (a.get("advertiser_visible"), a.get("copy_visible")) if x),
             "표시 URL": a.get("display_url_visible"), "링크": "미확보 (이미지 판독)", "판독 방식": (r.get("ai_read") or {}).get("method")}
            for r in by_kind("검색 화면") for a in (r.get("ai_read") or {}).get("ads", [])]
    tables = {
        "01_브랜드비교": result.get("comparison", []),
        "02_검색추이비교": trend,
        "03_시장검색어추이": market,
        "04_검색량합계": totals,
        "05_검색어별검색량": [clean(table_row(r, True)) for r in by_kind("검색량")],
        "06_연관검색어": [clean(table_row(r, True)) for r in by_kind("연관 검색어")],
        "20_뉴스핵심근거": result.get("news_digest", []),
        "07_뉴스": [clean(table_row(r, True)) | {"원문": r.get("text", "")} for r in by_kind("뉴스")],
        "08_검색화면광고관측": search_status_rows(records, None),
        "09_검색광고문구": ads,
        "10_Meta광고": [clean(table_row(r, True)) | {"전체 문구": r.get("text", "")} for r in by_kind("광고")],
        "11_UTM구조": [r for r in result.get("utm", []) if utm.has_values(r)],
        "12_브랜드별UTM구조": result.get("utm_structures") or utm.structure_rows(result.get("utm", [])),
        "13_SNS계정": [clean(table_row(r, True)) for r in by_kind("Instagram", "YouTube")],
        "14_SNS게시물": [clean(table_row(r, True)) | {"원문": r.get("text", "")} for r in by_kind("Instagram 게시물", "YouTube 게시물")],
        "15_공식페이지": [clean(table_row(r, True)) | {"원문": r["text"]} for r in by_kind("홈페이지")],
        "16_원본파일": [{"종류": r["kind"], "브랜드": r.get("brand"), "파일": a["filename"], "ZIP 경로": path, "역할": a.get("role"), "출처": a.get("source_url") or r["source_url"]} for r, a, path in asset_entries(records)],
        "17_수집조건": collection_info_v4(result),
        "18_수집상태": [{"대상": pt.get("label"), "상태": pt.get("state"), "안내": pt.get("message", "")} for pt in parts.values()],
        "19_이전수집대비": [{k: v for k, v in c.items() if k != "자료ID"} for c in result.get("changes", [])],
        "21_검색추이주요사실": all_trend_facts(parts),
        "22_뉴스주제묶음": [{"섹션": s["title"], "대표 제목": g["title"], "대표 발췌": g["excerpt"], "유사 기사 수": len(g["articles"]), "출처": "\n".join(r.get("source_url", "") for r in g["articles"])} for s in news_sections(by_kind("뉴스")) for g in s["groups"]],
        "23_맞춤통계자료": result.get("stat_resources", []),
        "24_검색관측전체상태": search_status_rows(records, None, observed_only=False),
    }
    return {k: v for k, v in tables.items() if v or k == "17_수집조건"}


def tables_for(result):
    return project_tables(result) if result.get("schema_version") in (4, 5) else report_tables(result)


def scalar(value):
    return json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else value


def excel(result):
    wb=Workbook()
    wb.remove(wb.active)
    for name,rows in tables_for(result).items():
        ws=wb.create_sheet(name)
        columns=list(dict.fromkeys(k for r in rows for k in r)) or ["안내"]
        ws.append(columns)
        for row in rows:
            values=[scalar(row.get(k)) for k in columns]
            # Preserve long text across continuation rows rather than silently truncating it.
            segments=max([1]+[(len(v)+29999)//30000 for v in values if isinstance(v,str)])
            for i in range(segments):
                ws.append([v[i*30000:(i+1)*30000] if isinstance(v,str) else v if i==0 else None for v in values])
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value,str):
                        cell.data_type="s"
                        if i==0 and cell.value.startswith(("https://","http://")) and len(cell.value)<=2000:
                            cell.hyperlink=cell.value
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
    data=BytesIO()
    wb.save(data)
    return data.getvalue()


def report_html(session,result):
    if result.get("schema_version") in (4, 5):
        from core.exporters.visual_report import report
        return report(session, result, project_tables(result))
    esc=lambda v:html.escape(str(scalar(v) if v is not None else "미제공"))
    def cell(value):
        if isinstance(value,str) and value.startswith(("https://","http://")):
            # 긴 URL은 표시 문자열만 줄이고 링크는 원문 그대로 둔다.
            text=value if len(value)<=70 else value[:67]+"…"
            return '<a href="'+html.escape(value,quote=True)+'" rel="noreferrer" title="'+html.escape(value,quote=True)+'">'+html.escape(text)+'</a>'
        return esc(value)
    sections=[]
    for title,rows in tables_for(result).items():
        columns=list(dict.fromkeys(k for r in rows for k in r))
        sections.append("<h2>"+esc(title)+"</h2><div class=scroll><table><tr>"+''.join("<th>"+esc(c)+"</th>" for c in columns)+"</tr>"+
                        ''.join("<tr>"+''.join("<td>"+cell(r.get(c))+"</td>" for c in columns)+"</tr>" for r in rows)+"</table></div>")
    included=0
    for row in result.get("records",[]):
        first=next((a for a in row.get("assets",[]) if a["filename"].endswith((".png",".jpg",".webp",".gif"))),None)
        data=read_bytes(first) if first else None
        if data is not None and included+len(data)<=20*1024**2:
            included+=len(data)
            mime="image/jpeg" if first["filename"].endswith(".jpg") else "image/"+first["filename"].rsplit('.',1)[1]
            sections.append('<figure><img style="max-width:100%" src="data:'+mime+';base64,'+base64.b64encode(data).decode()+'"><figcaption>'+esc(row["brand"])+" · "+esc(row["kind"])+" · "+cell(row["source_url"])+"</figcaption></figure>")
    # Flat report plus original file manifest: usable even when the source page disappears.
    return "<!doctype html><html lang=ko><meta charset=utf-8><meta name=viewport content='width=device-width,initial-scale=1'><title>ADetect 자료</title><style>body{font:15px system-ui;margin:32px;background:#101010;color:#eee}table{border-collapse:collapse}td,th{padding:10px;border-bottom:1px solid #555;text-align:left;white-space:pre-wrap;max-width:600px;overflow-wrap:anywhere}.scroll{overflow:auto}a{color:#ff8a65}.warn{border-left:3px solid #ff5a36;padding-left:10px}</style><h1>"+esc(session["brand_name"])+" · "+esc(session.get("campaign") or "브랜드 전체")+"</h1>"+("<p class=warn>날짜 혼합 — 선택한 자료의 수집일이 서로 다릅니다: "+esc(", ".join(result.get("inputs",{}).get("collection_dates",[])))+"</p>" if result.get("inputs",{}).get("date_mixed") else "")+"<p>"+("SAMPLE — " if result.get("sample_sources") else "")+"수집된 표본·원문 기록입니다. 브랜드의 주장을 검증된 사실이나 성과로 해석하지 않습니다.</p><p>검색지수는 검색어별 조회 기간 최고치=100. 누락 월은 0이 아닙니다.</p>"+''.join(sections)+"</html>"


def combined(session):
    from core.collection import all_records
    parts={stage+":"+key:value for stage in ("market","brand","creative") for key,value in session.get(stage+"_result",{}).get("parts",{}).items()}
    results=[session.get(s+"_result",{}) for s in ("market","brand","creative")]
    return {"schema_version":2,"status":"자료 모음","parts":parts,"records":all_records(session),
            "changes":[c for r in results for c in r.get("changes",[])],
            "inputs":{k:session.get(k) for k in ("brand_keywords","general_keywords","campaign","include_terms","exclude_terms")},
            "sample_sources":list({v for r in results for v in r.get("sample_sources",[])}),"summary":session.get("synthesis_result",{}).get("summary")}


def safe_component(value):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value or "공통"))[:70].strip(" .") or "공통"
    if name.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL", *[f"{p}{i}" for p in ("COM", "LPT") for i in range(1,10)]}:
        name = "_" + name
    return name


def asset_entries(records):
    names = sorted({str(r.get("brand") or "공통") for r in records})
    folders = {}
    for name in names:
        safe = safe_component(name)
        collisions = [n for n in names if safe_component(n).casefold() == safe.casefold()]
        folders[name] = safe + ("_" + hashlib.sha256(name.encode()).hexdigest()[:8] if len(collisions) > 1 else "")
    for row in records:
        for asset in row.get("assets", []):
            path = "originals/" + folders[str(row.get("brand") or "공통")] + "/" + safe_component(row.get("kind")) + "/" + safe_component(asset["filename"])
            yield row, asset, path


def package(session,result):
    output=BytesIO()
    omitted=[]
    with ZipFile(output,"w",ZIP_DEFLATED) as archive:
        raw=json.dumps(result,ensure_ascii=False,indent=2).encode()
        check_secrets(raw)
        archive.writestr("자료.json",raw)
        archive.writestr("자료.html",report_html(session,result))
        archive.writestr("자료.xlsx",excel(result))
        seen=set()
        total=0
        manifest = []
        for row, asset, path in asset_entries(result.get("records", [])):
            if path in seen: continue
            seen.add(path)
            data=read_bytes(asset)
            state = "포함"
            if data is None:
                state = "파일 없음 또는 무결성 확인 실패"
            elif total+len(data)>250*1024**2:
                state = "ZIP 250MB 원본 한도"
            else:
                total+=len(data)
                archive.writestr(path,data)
            if state != "포함": omitted.append(path + " — " + state)
            manifest.append({"브랜드": row.get("brand"), "종류": row.get("kind"), "파일": asset["filename"], "ZIP 경로": path, "상태": state})
        archive.writestr("원본_경로목록.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        archive.writestr("원본_보관안내.txt","\n".join(omitted) or "자료 목록의 저장된 원본 파일을 모두 포함했습니다. 원격 링크만 있는 파일은 포함되지 않습니다.")
    return output.getvalue()
